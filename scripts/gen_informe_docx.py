"""Genera InformeFinal_DFSha.docx a partir de la estructura del Hito 1."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "InformeFinal_DFSha.docx"

NAVY = (0x1A, 0x36, 0x5D)


def set_run_font(run, size=11, bold=False, color=None, italic=False, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 14 if level == 1 else 12, bold=True, color=NAVY)
    return p


def add_p(doc, text, bold=False, italic=False, size=11, center=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size, bold=bold, italic=italic)
    return p


def shade_header(cell, color="1A365D"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size, bold=bold, color=color)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255), size=9)
        shade_header(t.rows[0].cells[i])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            set_cell_text(t.rows[r_i + 1].cells[c_i], val, size=9)
            if r_i % 2 == 1:
                shade_header(t.rows[r_i + 1].cells[c_i], "E8EEF6")
    doc.add_paragraph()
    return t


def add_mono(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size, font="Consolas")
    return p


def main() -> None:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    add_p(doc, "Universidad EAFIT", bold=True, size=16, center=True, space_after=2)
    add_p(
        doc,
        "ST0263 — Tópicos Especiales en Telemática / SI3007 Sistemas Distribuidos",
        size=11,
        center=True,
        space_after=12,
    )
    add_p(doc, "Proyecto 1: DFSha", bold=True, size=22, center=True, space_after=2)
    add_p(
        doc,
        "Sistema de Archivos Distribuido con Alta Disponibilidad",
        size=13,
        center=True,
        space_after=10,
    )
    add_p(doc, "Informe final — Hitos 1, 2 y 3", bold=True, size=14, center=True, space_after=4)
    add_p(
        doc,
        "Especificación · Arquitectura y comunicaciones · HA, réplica, consistencia y seguridad",
        italic=True,
        size=11,
        center=True,
        space_after=14,
    )
    add_p(doc, "2026-2", size=11, center=True, space_after=2)
    add_p(doc, "Equipo: Carlos · James · David · Felipe", size=11, center=True, space_after=2)
    add_p(doc, "https://github.com/FelipeOchoaL/DFSha", size=10, center=True, space_after=18)

    add_heading_custom(doc, "1. Objetivo general")
    add_p(
        doc,
        "Diseñar e implementar DFSha, un sistema de archivos distribuido por bloques "
        "(Opción 1 del enunciado: C/S + S2S) que permita almacenar y recuperar archivos "
        "grandes con transparencia para el usuario, sin que el maestro transporte bytes. "
        "El sistema garantiza partición, replicación factor 3, consistencia por checksum, "
        "alta disponibilidad del NameNode y un camino de seguridad (TLS / JWT) activable.",
    )

    add_heading_custom(doc, "2. Descripción del servicio")
    add_p(
        doc,
        "DFSha sigue la línea HDFS / GFS: un maestro de metadatos y trabajadores de datos. "
        "No es POSIX (fuera de alcance el RF3 del enunciado). El usuario ve rutas tipo "
        "/demo/archivo.bin; no ve IPs ni números de bloque.",
    )
    add_heading_custom(doc, "2.1 Cliente DFSha (Python 3.11+)", 2)
    add_p(
        doc,
        "CLI local (python -m client.main). Habla REST con el NameNode (metadatos) y gRPC "
        "directo con los DataNodes (bytes). Comandos: login, put, get, ls, cd, pwd, mkdir, "
        "rmdir, rm. Si el activo responde 503 o no responde, prueba el siguiente URL de "
        "DFSHA_NAMENODES.",
    )
    add_heading_custom(doc, "2.2 Capa de control — NameNode (Go 1.22+)", 2)
    add_p(
        doc,
        "Solo metadatos: namespace, tabla de bloques, colocación, heartbeats y edit log. "
        "Expone REST en :8080 y gRPC en :9000. Se despliega activo/standby: el standby "
        "replica el edit log por gRPC y se promueve si el activo deja de responder a Ping.",
    )
    add_heading_custom(doc, "2.3 Capa de almacenamiento — DataNodes (Go 1.22+)", 2)
    add_p(
        doc,
        "Almacenan bloques en disco, se registran, envían heartbeat y BlockReceived, y "
        "ejecutan el pipeline de réplica (WriteBlock hacia el siguiente hop). Cuatro nodos "
        "en racks A/A/B/B.",
    )
    add_p(
        doc,
        "Qué no se usó (y por qué). El enunciado sugiere FastAPI, Kafka o una BD como "
        "ejemplos. DFSha no los necesita: el NameNode es net/http en Go; la coordinación "
        "es heartbeat + comandos en la respuesta (no cola MOM); la persistencia de "
        "metadatos es edit log + snapshot en disco. Menos piezas, mismo contrato.",
    )

    add_heading_custom(doc, "3. Requerimientos funcionales")
    add_table(
        doc,
        ["ID", "Operación", "Estado", "Cómo quedó"],
        [
            [
                "RF1",
                "put <archivo>",
                "Hecho",
                "El cliente parte el archivo en bloques de 8 MB; el NameNode asigna un pipeline de 3 DataNodes; el cliente escribe solo al primero.",
            ],
            [
                "RF2",
                "get <archivo>",
                "Hecho",
                "GET /files/locations + ReadBlock en paralelo; si una réplica falla o el SHA-256 no coincide, se prueba la siguiente.",
            ],
            [
                "RF3*",
                "ls [ruta]",
                "Hecho",
                "Namespace virtual en el NameNode. El RF3 POSIX del enunciado no se implementó.",
            ],
            ["RF4", "mkdir <ruta>", "Hecho", "POST /fs/mkdir."],
            [
                "RF5",
                "rm <archivo>",
                "Hecho",
                "Borra del namespace y encola DeleteCommand por heartbeat.",
            ],
            ["RF6", "rmdir <ruta>", "Hecho", "Solo directorio vacío."],
            [
                "RF7",
                "Autenticación",
                "Parcial",
                "El CLI tiene login y sesión JWT. El NameNode tiene Identify / ACL; en la demo las peticiones van como usuario anónimo.",
            ],
        ],
    )

    add_heading_custom(doc, "4. Requerimientos no funcionales")
    add_p(doc, "Cada RNF del Hito 1 y cómo se resolvió al cierre:")
    add_table(
        doc,
        ["ID", "Atributo", "Hito 1", "Implementación final"],
        [
            [
                "RNF1",
                "Escalabilidad",
                "Bloques 64 MB; DN en caliente",
                "Bloques de 8 MB. Un DataNode nuevo se registra por gRPC sin reiniciar. Se bajó de 64 a 8 MB para t2.small y demos cortas.",
            ],
            [
                "RNF2",
                "Alta disponibilidad",
                "Activo/standby + RF=3",
                "Edit log replicado; Ping + promoción; el antiguo activo se degrada. Réplica 3 en ≥2 racks. Re-replicación si un DN pierde 3 heartbeats.",
            ],
            [
                "RNF3",
                "Consistencia",
                "Quorum 2/3",
                "Más estricto: commit espera 3 BlockReceived con el mismo SHA-256 (máx. 3 s). Si no, 409 y otro pipeline.",
            ],
            [
                "RNF4",
                "Particionamiento",
                "64 MB + racks",
                "Bloques consecutivos no comparten primario; réplicas cubren otro rack antes de repetir uno.",
            ],
            [
                "RNF5",
                "Rendimiento",
                "Cliente↔DN; pipeline",
                "El NameNode no ve el payload. Pipeline DN1→DN2→DN3. get en paralelo.",
            ],
            [
                "RNF6",
                "Seguridad",
                "HTTPS, mTLS, JWT, AES-256",
                "TLS/mTLS listo en código. Demo: TLS_ENABLED=false. JWT y AES-GCM diseñados; no son el camino de la demo.",
            ],
            [
                "RNF7",
                "Transparencia",
                "CLI tipo Linux",
                "El usuario solo usa rutas. Failover y reintento de réplica son internos.",
            ],
        ],
    )

    add_heading_custom(doc, "5. Arquitectura (Hito 2)")
    add_p(
        doc,
        "Arquitectura maestro–trabajadores en red privada AWS Academy (o docker compose "
        "local). El NameNode no transporta bytes.",
    )
    add_mono(
        doc,
        "CLI Python\n"
        "   │  REST  (dónde guardar / dónde leer)\n"
        "   ▼\n"
        "NameNode activo  ←── gRPC edit log + Ping ──→  NameNode standby\n"
        "    │  gRPC (Register, Heartbeat, BlockReport, BlockReceived)\n"
        "    ▼\n"
        "DataNode 1 (rack A) ──pipeline WriteBlock──→ DataNode 2 (rack A) ──→ DataNode 3 (rack B)",
    )
    add_table(
        doc,
        ["Componente", "Tecnología", "Rol"],
        [
            ["Cliente", "Python 3.11+", "CLI. REST al NameNode, gRPC a DataNodes."],
            [
                "NameNode activo",
                "Go 1.22+",
                "Namespace, colocación, commit 3/3, REST :8080, gRPC :9000.",
            ],
            [
                "NameNode standby",
                "Mismo binario, NAMENODE_ROLE=standby",
                "Réplica del edit log. REST responde 503 hasta promoverse.",
            ],
            [
                "DataNode ×4",
                "Go 1.22+",
                "Disco + pipeline + heartbeat. gRPC :50051–:50054.",
            ],
        ],
    )
    add_p(
        doc,
        "Contratos: proto/datanode.proto, proto/namenode.proto. REST: documentation/api-rest.md.",
    )

    add_heading_custom(doc, "5.1 Infraestructura AWS", 2)
    add_table(
        doc,
        ["Instancia", "Rol", "IP privada", "Puertos"],
        [
            ["dfsha-nn-active", "NameNode activo", "172.31.0.10", "8080 REST, 9000 gRPC"],
            ["dfsha-nn-standby", "NameNode standby", "172.31.0.11", "8080 REST, 9000 gRPC"],
            ["dfsha-dn-1 … dn-4", "DataNodes, racks A/A/B/B", "172.31.0.21–.24", "50051 gRPC"],
        ],
    )
    add_p(
        doc,
        "El DataNode publica ADVERTISE_ADDR = IP pública: es la que el NameNode entrega "
        "al CLI que corre en el PC del usuario. Guía: documentation/guia-aws.md.",
    )

    add_heading_custom(doc, "6. Protocolos de comunicación (Hito 2)")
    add_table(
        doc,
        ["Conexión", "Protocolo", "Qué viaja"],
        [
            [
                "Cliente ↔ NameNode",
                "REST / JSON",
                "POST /auth/login, /fs/ls|mkdir, /fs/rm, /fs/rmdir, /files/allocate, /files/commit, /files/finalize, GET /files/locations.",
            ],
            [
                "Cliente ↔ DataNode",
                "gRPC streaming",
                "WriteBlock (header + chunks + SHA-256 + pipeline), ReadBlock.",
            ],
            ["NameNode ↔ NameNode", "gRPC", "ReplicateEditLog, Ping."],
            [
                "DataNode ↔ NameNode",
                "gRPC",
                "Register, Heartbeat (~3 s, trae Replicate/Delete), BlockReport, BlockReceived.",
            ],
            [
                "DataNode ↔ DataNode",
                "gRPC",
                "Mismo WriteBlock: el hop anterior reenvía el bloque al siguiente.",
            ],
        ],
    )
    add_p(doc, "No hay bus de mensajes. El heartbeat es el canal de control asíncrono.")

    add_heading_custom(doc, "7. Algoritmo de particionamiento y distribución")
    add_heading_custom(doc, "7.1 División", 2)
    add_p(
        doc,
        "El cliente corta el archivo en bloques de 8 MB (el último puede ser menor). "
        "El NameNode asigna el blockId. Un archivo de 20 MB son 3 bloques en DataNodes distintos.",
    )
    add_heading_custom(doc, "7.2 Colocación (placement.go)", 2)
    add_p(
        doc,
        "Candidatos = DataNodes vivos con espacio. Si hay menos que el factor de réplica → 503.",
    )
    add_p(
        doc,
        "1. Primario: mayor espacio libre, excluyendo el primario del bloque anterior del mismo archivo.",
    )
    add_p(
        doc,
        "2. Réplicas 2 y 3: se elige el nodo que minimiza réplicas ya puestas en su rack "
        "(y, empate, más espacio libre). Con dos racks esto cubre otro rack antes de repetir uno.",
    )
    add_p(
        doc,
        "Así un fallo de rack no pierde el archivo y ningún DataNode tiene todos los bloques.",
    )
    add_heading_custom(doc, "7.3 Escritura (put)", 2)
    add_p(doc, "1. POST /files/allocate → blockId + [DN1, DN2, DN3].")
    add_p(doc, "2. Cliente → DN1: WriteBlock (bytes + SHA-256 + pipeline [DN2, DN3]).")
    add_p(doc, "3. DN1 persiste, reenvía a DN2; DN2 a DN3. El ACK sube por la cadena.")
    add_p(doc, "4. Cada DN avisa BlockReceived.")
    add_p(doc, "5. POST /files/commit espera 3 confirmaciones con el mismo checksum (máx. 3 s).")
    add_p(doc, "6. Tras el último bloque, POST /files/finalize. Cada mutación se replica al standby.")
    add_heading_custom(doc, "7.4 Lectura (get)", 2)
    add_p(doc, "1. GET /files/locations → bloques ordenados y réplicas vivas.")
    add_p(doc, "2. Descarga gRPC en paralelo; checksum; siguiente réplica si falla.")
    add_p(doc, "3. Concatenación en el cliente. El NameNode no interviene en la transferencia.")

    add_heading_custom(doc, "8. Alta disponibilidad, replicación y consistencia (Hito 3)")
    add_p(
        doc,
        "Replicación. Factor 3, pipeline, rack awareness. Si un DataNode pierde 3 heartbeats "
        "se marca muerto; sus bloques quedan under-replicated y el NameNode mete "
        "ReplicateCommand en el heartbeat de un nodo que sí los tiene.",
    )
    add_p(
        doc,
        "Consistencia. Write-once. Un bloque no entra al namespace hasta 3 réplicas con el "
        "mismo SHA-256. No hay quorum 2/3: un put incompleto no se publica. En lectura, "
        "checksum mismatch descarta esa copia.",
    )
    add_p(
        doc,
        "HA del NameNode. El activo envía cada entrada del edit log al standby "
        "(HA_REQUIRE_ACK=true en producción: no confirma al cliente si el standby no aplicó; "
        "en compose local va false para no bloquear el arranque). El standby hace Ping; "
        "tras N fallos se promueve, pide re-registro y espera block reports. El CLI ya "
        "trae las dos URLs y conmuta solo.",
    )

    add_heading_custom(doc, "9. Seguridad (Hito 3)")
    add_table(
        doc,
        ["Aspecto", "Diseño", "Estado al cierre"],
        [
            [
                "Autenticación",
                "Usuario/clave, JWT 24 h, ACL por dueño",
                "CLI: login + ~/.dfsha/session.json. NameNode: Identify y permisos listos; demo sin middleware → usuario anónimo.",
            ],
            [
                "En tránsito",
                "HTTPS en REST, mTLS gRPC (TLS ≥ 1.2)",
                "Implementado y cubierto por tests. Demo: TLS_ENABLED=false. Activación: deploy/certs/gen_certs.sh + TLS_ENABLED=true.",
            ],
            [
                "En reposo",
                "AES-256-GCM por bloque",
                "Gancho en el DataNode; no activo en la demo.",
            ],
            [
                "Entre nodos",
                "CA del proyecto",
                "Igual que TLS: código listo, apagado en el primer despliegue.",
            ],
        ],
    )
    add_p(
        doc,
        "La demo académica prioriza put / get / HA / réplica. Encender TLS es un cambio "
        "de variables, no de arquitectura.",
    )

    add_heading_custom(doc, "10. Entorno")
    add_p(
        doc,
        "Local. docker compose up --build → NN activo :8080, standby :8081, DN 50051–50054. "
        "El CLI usa host.docker.internal.",
    )
    add_mono(
        doc,
        "python -m client.main mkdir /demo\n"
        "python -m client.main put .\\archivo.bin /demo/archivo.bin\n"
        "python -m client.main get /demo/archivo.bin .\\copia.bin",
        size=10,
    )
    add_p(
        doc,
        "AWS. Tras git pull en cada EC2: copiar deploy/aws/env/<nodo>.env.example a "
        "/opt/dfsha/node.env y correr namenode.sh o datanode.sh.",
    )
    add_p(
        doc,
        "Pruebas. cd services && go test ./...  ·  python -m pytest client/tests -q "
        "(pruebas del CLI con NameNode/DataNodes falsos).",
    )

    add_heading_custom(doc, "11. Cumplimiento por hito")
    add_table(
        doc,
        ["Hito", "Pedido", "Entrega"],
        [
            [
                "1",
                "Especificación definitiva",
                "Este documento actualiza el Hito 1: 8 MB, consistencia 3/3, 4 DN, 6 EC2, sin Kafka/FastAPI/BD.",
            ],
            [
                "2",
                "Arquitectura y comunicaciones",
                "NameNode REST + HA gRPC, DataNode I/O y pipeline, cliente Python, protos, Docker, scripts AWS. put/get/ls/mkdir/rm funcionales.",
            ],
            [
                "3",
                "HA, réplica, consistencia, seguridad",
                "Activo/standby con edit log y failover; réplica 3 + re-replicación; commit 3/3 SHA-256; TLS/JWT en código, TLS off en la demo.",
            ],
        ],
    )
    add_p(doc, "Fuera de alcance: semántica POSIX (append, locks, seek).")

    add_heading_custom(doc, "12. Referencias")
    for i, ref in enumerate(
        [
            "Ghemawat, S., Gobioff, H. y Leung, S. T. (2003). The Google File System. ACM SOSP.",
            "Shvachko, K. et al. (2010). The Hadoop Distributed File System. IEEE MSST.",
            "gRPC. https://grpc.io/docs/",
            "Enunciado ST0263 / SI3007 — Proyecto 1 DFS, 2026-2.",
            "Código y contratos: repositorio DFSha (proto/, services/, client/).",
        ],
        1,
    ):
        add_p(doc, f"{i}. {ref}", size=10, space_after=3)

    doc.save(str(OUT))
    print(f"saved {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
